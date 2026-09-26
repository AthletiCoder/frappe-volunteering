# Copyright (c) 2026, Vadiraj Tirtha Das and contributors
# For license information, please see license.txt

import base64
from decimal import Decimal
from io import BytesIO
from unittest.mock import patch

import frappe
from frappe.tests import IntegrationTestCase, UnitTestCase

from volunteering.volunteering import invoice_generator as invoices
from volunteering.volunteering.invoice_generator import (
	_build_docx,
	_normalise_payload,
	_render_pdf_html,
	_safe_filename,
)


def _payload(invoice_type="GST"):
	return {
		"invoice_type": invoice_type,
		"invoice_number": "INV/2026/001",
		"invoice_date": "2026-09-16",
		"supplier": {
			"name": "Example Kitchen Supplies",
			"address": "12 Market Road, Mumbai",
			"state": "Maharashtra",
			"pin_code": "400001",
			"gstin": "27ABCDE1234F1Z5",
			"pan": "ABCDE1234F",
		},
		"consignee": {
			"name": "Sevamrita Foundation",
			"address": "1 Service Road, Mumbai",
			"state": "Maharashtra",
			"pin_code": "400002",
			"gstin": "27AAECS1234F1Z5",
		},
		"buyer_same_as_consignee": True,
		"items": [
			{
				"description": "Steel utensils",
				"hsn_sac": "7323",
				"unit": "Nos",
				"quantity": 2,
				"rate": 100,
			},
		],
		"transportation_charges": 10,
		"other_charges": 0,
		"gst_amount": 37.80,
		"volunteer": {"name": "Test Volunteer", "employee": "HR-EMP-TEST"},
		"volunteer_signature_data": _signature_data(),
		"vendor_will_sign": False,
	}


def _signature_data(blank=False, transparent=False):
	from PIL import Image, ImageDraw

	image = Image.new(
		"RGBA" if transparent else "RGB", (600, 180), (255, 255, 255, 0) if transparent else "white"
	)
	if not blank:
		draw = ImageDraw.Draw(image)
		draw.line([(80, 120), (180, 55), (260, 130), (390, 45), (510, 110)], fill="black", width=6)
	buffer = BytesIO()
	image.save(buffer, format="PNG")
	return "data:image/png;base64," + base64.b64encode(buffer.getvalue()).decode()


class UnitTestInvoiceGenerator(UnitTestCase):
	def test_generation_retries_transaction_conflict(self):
		result = {"invoice_number": "INV-2098-000001"}
		with (
			patch.object(
				invoices,
				"_generate_invoice_documents",
				side_effect=[frappe.QueryDeadlockError("conflict"), result],
			) as generate,
			patch.object(frappe.db, "rollback") as rollback,
		):
			self.assertEqual(invoices.generate_invoice_documents({}), result)
		self.assertEqual(generate.call_count, 2)
		rollback.assert_called_once_with()

	def test_generation_transaction_retries_are_bounded(self):
		with (
			patch.object(
				invoices, "_generate_invoice_documents", side_effect=frappe.QueryDeadlockError("conflict")
			) as generate,
			patch.object(frappe.db, "rollback") as rollback,
		):
			with self.assertRaisesRegex(frappe.ValidationError, "generation is busy"):
				invoices.generate_invoice_documents({})
		self.assertEqual(generate.call_count, 3)
		self.assertEqual(rollback.call_count, 3)

	def test_gst_totals_are_recalculated_on_server(self):
		data = _normalise_payload(_payload())
		self.assertEqual(data["items_total"], Decimal("200.00"))
		self.assertEqual(data["taxable_total"], Decimal("210.00"))
		self.assertEqual(data["gst_amount"], Decimal("37.80"))
		self.assertEqual(data["grand_total"], Decimal("247.80"))

	def test_buyer_is_always_the_selected_consignee(self):
		payload = _payload()
		payload["buyer_same_as_consignee"] = False
		payload["buyer"] = {
			"name": "Forged Buyer",
			"address": "Another address",
			"state": "Goa",
		}
		data = _normalise_payload(payload)
		self.assertTrue(data["buyer_same_as_consignee"])
		self.assertEqual(data["buyer"], data["consignee"])
		self.assertNotEqual(data["buyer"]["name"], "Forged Buyer")

	def test_mandatory_volunteer_signature_is_validated_and_embedded_in_pdf_and_word(self):
		from docx import Document

		data = _normalise_payload(_payload())
		self.assertTrue(data["volunteer_signature_png"].startswith(b"\x89PNG"))
		self.assertIn("data:image/png;base64,", _render_pdf_html(data))
		document = Document(BytesIO(_build_docx(data)))
		self.assertEqual(len(document.inline_shapes), 1)
		word_text = "\n".join(
			cell.text for table in document.tables for row in table.rows for cell in row.cells
		)
		for text in (_render_pdf_html(data), word_text):
			self.assertIn("Volunteer reimbursement declaration", text)
			self.assertIn("I confirm that I paid the amount shown above", text)
			self.assertIn("request reimbursement to my bank account", text)

	def test_missing_blank_or_non_png_volunteer_signature_is_rejected(self):
		for signature in (
			"",
			_signature_data(blank=True),
			_signature_data(blank=True, transparent=True),
			"data:image/png;base64,Zm9yZ2Vk",
			"not-an-image",
		):
			with self.subTest(signature=signature[:30]):
				payload = _payload()
				payload["volunteer_signature_data"] = signature
				with self.assertRaises(frappe.ValidationError):
					_normalise_payload(payload)

	def test_non_gst_without_vendor_signature_has_no_vendor_declaration(self):
		payload = _payload("NON_GST")
		data = _normalise_payload(payload)
		self.assertEqual(data["gst_amount"], Decimal("0.00"))
		self.assertFalse(data["vendor_will_sign"])
		self.assertEqual(data["vendor_declaration"], "")
		self.assertNotIn("Vendor confirmation", _render_pdf_html(data))

	def test_non_gst_pan_is_optional_but_checked_when_supplied(self):
		payload = _payload("NON_GST")
		payload["supplier"].pop("pan")
		self.assertEqual(_normalise_payload(payload)["supplier"]["pan"], "")
		payload["supplier"]["pan"] = "INVALID"
		with self.assertRaises(frappe.ValidationError):
			_normalise_payload(payload)

	def test_hsn_sac_is_optional_for_both_invoice_types(self):
		for invoice_type in ("NON_GST", "GST"):
			for code in (None, "", "   ", "7323"):
				with self.subTest(invoice_type=invoice_type, code=code):
					payload = _payload(invoice_type)
					if code is None:
						payload["items"][0].pop("hsn_sac")
					else:
						payload["items"][0]["hsn_sac"] = code
					data = _normalise_payload(payload)
					self.assertEqual(data["items"][0]["hsn_sac"], "7323" if code == "7323" else "")

	def test_missing_volunteer_identity_is_rejected(self):
		payload = _payload()
		payload.pop("volunteer")
		with self.assertRaises(frappe.ValidationError):
			_normalise_payload(payload)

	def test_server_volunteer_identity_overrides_browser_and_vendor_is_absent_by_default(self):
		from docx import Document

		for invoice_type in ("NON_GST", "GST"):
			with self.subTest(invoice_type=invoice_type):
				payload = _payload(invoice_type)
				payload["volunteer"] = {"name": "Forged Volunteer", "employee": "Forged Employee"}
				data = _normalise_payload(
					payload, volunteer_override={"name": "Test Volunteer", "employee": "HR-EMP-TEST"}
				)
				self.assertEqual(data["title"], "TAX INVOICE" if invoice_type == "GST" else "INVOICE")
				document = Document(BytesIO(_build_docx(data)))
				word_text = "\n".join(
					[p.text for p in document.paragraphs]
					+ [cell.text for table in document.tables for row in table.rows for cell in row.cells]
				)
				for text in (_render_pdf_html(data), word_text):
					self.assertIn("Volunteer reimbursement declaration", text)
					self.assertIn("Volunteer Signature", text)
					self.assertIn("Test Volunteer", text)
					self.assertIn("HR-EMP-TEST", text)
					self.assertNotIn("Forged", text)
					self.assertNotIn("Vendor confirmation", text)
					self.assertNotIn("not registered", text)
					self.assertNotIn("Non-GST declaration", text)
					self.assertNotIn("Authorised signatory and supplier signature", text)
					self.assertNotIn("EXPENSE STATEMENT", text)
					self.assertNotIn("Prepared for volunteer", text)

	def test_vendor_signature_and_non_gst_declaration_are_conditional_in_both_documents(self):
		from docx import Document

		for invoice_type in ("NON_GST", "GST"):
			for signatory in ("Asha Vendor", ""):
				with self.subTest(invoice_type=invoice_type, signatory=signatory):
					payload = _payload(invoice_type)
					payload["vendor_will_sign"] = True
					payload["vendor_signature_data"] = _signature_data()
					payload["authorised_signatory"] = signatory
					data = _normalise_payload(payload)
					doc = Document(BytesIO(_build_docx(data)))
					word_text = "\n".join(
						[p.text for p in doc.paragraphs]
						+ [cell.text for table in doc.tables for row in table.rows for cell in row.cells]
					)
					for text in (_render_pdf_html(data), word_text):
						self.assertIn("Vendor confirmation", text)
						if invoice_type == "NON_GST":
							self.assertIn("not registered", text)
							self.assertIn("does not have a GSTIN", text)
							self.assertIn(signatory or "the undersigned", text)
						else:
							self.assertNotIn("Non-GST declaration", text)
							self.assertNotIn("not registered", text)
						self.assertIn("Authorised Signatory", text)
						self.assertIn("Volunteer reimbursement declaration", text)
					self.assertEqual(len(doc.inline_shapes), 2)

	def test_vendor_signing_requires_a_vendor_signature(self):
		payload = _payload()
		payload["vendor_will_sign"] = True
		with self.assertRaisesRegex(frappe.ValidationError, "vendor to sign"):
			_normalise_payload(payload)

	def test_gst_invoice_number_is_limited_to_sixteen_characters(self):
		payload = _payload()
		payload["invoice_number"] = "GST-INVOICE-TOO-LONG"
		with self.assertRaises(frappe.ValidationError):
			_normalise_payload(payload)

	def test_generated_number_override_does_not_require_browser_number(self):
		payload = _payload()
		payload.pop("invoice_number")
		data = _normalise_payload(payload, invoice_number_override="INV-2098-000001")
		self.assertEqual(data["invoice_number"], "INV-2098-000001")

	def test_generated_number_override_ignores_browser_number(self):
		payload = _payload()
		payload["invoice_number"] = "FORGED-NUMBER-THAT-IS-TOO-LONG-FOR-GST"
		data = _normalise_payload(payload, invoice_number_override="INV-2098-000001")
		self.assertEqual(data["invoice_number"], "INV-2098-000001")

	def test_non_finite_amount_is_rejected(self):
		payload = _payload()
		payload["items"][0]["rate"] = "NaN"
		with self.assertRaises(frappe.ValidationError):
			_normalise_payload(payload)

	def test_html_escapes_supplier_content(self):
		payload = _payload()
		payload["supplier"]["name"] = "Vendor <script>alert(1)</script>"
		html = _render_pdf_html(_normalise_payload(payload))
		self.assertNotIn("<script>", html)
		self.assertIn("&lt;script&gt;", html)

	def test_docx_contains_invoice_identity_and_total(self):
		from docx import Document

		content = _build_docx(_normalise_payload(_payload()))
		document = Document(BytesIO(content))
		text = "\n".join(
			[paragraph.text for paragraph in document.paragraphs]
			+ [cell.text for table in document.tables for row in table.rows for cell in row.cells]
		)
		self.assertIn("TAX INVOICE", text)
		self.assertIn("INV/2026/001", text)
		self.assertIn("247.80", text)
		self.assertNotIn("Remittance Details", text)
		self.assertNotIn("expense reimbursement", text.lower())
		self.assertNotIn("Receipt review", text)

	def test_supplier_bank_is_optional_and_omitted_from_pdf_and_word_when_blank(self):
		from docx import Document

		data = _normalise_payload(_payload())
		self.assertFalse(invoices._has_bank(data["bank"]))
		self.assertNotIn("Remittance Details", _render_pdf_html(data))
		document = Document(BytesIO(_build_docx(data)))
		word_text = "\n".join(
			cell.text for table in document.tables for row in table.rows for cell in row.cells
		)
		self.assertNotIn("Remittance Details", word_text)

	def test_supplier_bank_is_rendered_without_employee_reimbursement_language(self):
		from docx import Document

		payload = _payload()
		payload["bank"] = {
			"account_name": "Example Kitchen Supplies",
			"bank_name": "Vendor Cooperative Bank",
			"account_number": "123456789012",
			"ifsc": "VEND0123456",
			"branch": "Market Road",
			"swift": "VENDINBB",
			"upi_id": "example.vendor@upi",
		}
		data = _normalise_payload(payload)
		html = _render_pdf_html(data)
		document = Document(BytesIO(_build_docx(data)))
		word_text = "\n".join(
			cell.text for table in document.tables for row in table.rows for cell in row.cells
		)
		for content in (html, word_text):
			self.assertIn("Supplier Remittance Details", content)
			self.assertIn("Vendor Cooperative Bank", content)
			self.assertIn("123456789012", content)
			self.assertIn("example.vendor@upi", content)
			self.assertNotIn("employee reimbursement", content.lower())

	def test_twenty_styles_are_used_before_repeating(self):
		assignments = []
		chosen = []
		for index in range(21):
			style = invoices._choose_invoice_style(assignments)
			chosen.append(style)
			assignments.append({"vendor_key": f"vendor-{index}", "invoice_style": style})
		self.assertEqual(len(invoices.INVOICE_STYLES), 20)
		self.assertEqual(len(set(chosen[:20])), 20)
		self.assertEqual(chosen[20], invoices.INVOICE_STYLES[0]["id"])
		self.assertEqual(len({style["name"] for style in invoices.INVOICE_STYLES}), 20)
		self.assertEqual(
			invoices._choose_invoice_style(
				[
					{"vendor_name": "Same Vendor", "invoice_style": "style-01"},
					{"vendor_name": "same-vendor", "invoice_style": "style-02"},
				]
			),
			"style-02",
		)

	def test_twenty_styles_render_distinct_documents(self):
		data = _normalise_payload(_payload())
		rendered = set()
		for style in invoices.INVOICE_STYLES:
			data["invoice_style"] = style["id"]
			rendered.add(_render_pdf_html(data))
		self.assertEqual(len(rendered), 20)

	def test_vendor_style_identity_survives_address_gstin_and_bank_changes(self):
		first = {
			"name": " Example Kitchen Supplies Pvt. Ltd. ",
			"address": "First address",
			"state": "Maharashtra",
			"pin_code": "411001",
			"gstin": "",
		}
		changed = {
			"name": "example kitchen supplies pvt ltd",
			"address": "A completely new address",
			"state": "Telangana",
			"pin_code": "500001",
			"gstin": "36ABCDE1234F1Z2",
		}
		self.assertEqual(invoices._vendor_identity_key(first), invoices._vendor_identity_key(changed))

	def test_safe_filename_removes_path_characters(self):
		self.assertEqual(_safe_filename("../../INV 001"), "invoice-INV-001")


class IntegrationTestGeneratedInvoiceNumbers(IntegrationTestCase):
	def setUp(self):
		super().setUp()
		frappe.db.savepoint("generated_invoice_test")
		# Use future years under a rollback-only test transaction; existing counters
		# (if any) are restored, and the live generation year's counter is untouched.
		for year in (2098, 2099):
			frappe.db.sql(
				"UPDATE `tabSeries` SET `current`=0 WHERE `name`=%s",
				(f"SEVAMRITA-FORM-INVOICE-{year}-",),
			)

	def tearDown(self):
		frappe.db.rollback(save_point="generated_invoice_test")
		super().tearDown()

	def generation_context(self, pdf=None):
		from contextlib import ExitStack

		stack = ExitStack()
		stack.enter_context(patch.object(invoices, "nowdate", return_value="2098-12-31"))
		stack.enter_context(patch.object(invoices, "_require_employee", return_value="Test Employee"))
		stack.enter_context(patch.object(invoices, "_remember_vendor_address", return_value={}))
		stack.enter_context(patch.object(invoices, "_save_employee_signature"))
		stack.enter_context(
			patch.object(invoices, "_vendor_invoice_style", return_value=("vendor-key", "style-01"))
		)
		stack.enter_context(
			patch.object(
				invoices,
				"_employee_signer",
				side_effect=lambda employee: {"employee": employee, "name": "Test Volunteer"},
			)
		)
		stack.enter_context(
			patch.object(
				invoices,
				"_build_pdf",
				side_effect=pdf or (lambda data: invoices._render_pdf_html(data).encode()),
			)
		)
		return stack

	def test_sequence_is_shared_by_employee_and_invoice_type_and_generation_year(self):
		with self.generation_context():
			first = invoices.generate_invoice_documents(_payload("NON_GST"))
			with patch.object(invoices, "_require_employee", return_value="Another Employee"):
				second = invoices.generate_invoice_documents(_payload("GST"))
		self.assertEqual(first["invoice_number"], "INV-2098-000001")
		self.assertEqual(second["invoice_number"], "INV-2098-000002")

	def test_new_calendar_generation_year_starts_new_sequence(self):
		with patch.object(invoices, "nowdate", return_value="2098-12-31"):
			self.assertEqual(invoices._next_invoice_number(), "INV-2098-000001")
		with patch.object(invoices, "nowdate", return_value="2099-01-01"):
			self.assertEqual(invoices._next_invoice_number(), "INV-2099-000001")
		with patch.object(invoices, "nowdate", return_value="2098-12-31"):
			self.assertEqual(invoices._next_invoice_number(), "INV-2098-000002")

	def test_pdf_only_does_not_generate_word_and_word_reuses_number(self):
		payload = _payload("NON_GST")
		payload["supplier"].pop("pan")
		with self.generation_context(), patch.object(invoices, "_build_docx") as word:
			pdf = invoices.generate_invoice_documents(payload, output_format="pdf")
			word.assert_not_called()
			self.assertIn("pdf", pdf)
			self.assertNotIn("docx", pdf)
		with self.generation_context(), patch.object(invoices, "_build_pdf") as build_pdf:
			word = invoices.generate_invoice_documents(
				payload, output_format="docx", generation_reference=pdf["generation_reference"]
			)
			build_pdf.assert_not_called()
		self.assertIn("docx", word)
		self.assertNotIn("pdf", word)
		self.assertEqual(pdf["invoice_number"], word["invoice_number"])
		with self.generation_context():
			self.assertEqual(invoices._next_invoice_number(), "INV-2098-000002")

	def test_word_can_be_generated_first_without_pdf(self):
		with self.generation_context(), patch.object(invoices, "_build_pdf") as pdf:
			word = invoices.generate_invoice_documents(_payload(), output_format="docx")
			pdf.assert_not_called()
		self.assertIn("docx", word)
		self.assertNotIn("pdf", word)

	def test_volunteer_identity_is_server_assigned_and_adding_vendor_signature_gets_new_number(self):
		payload = _payload("NON_GST")
		with self.generation_context():
			volunteer_only = invoices.generate_invoice_documents(payload, output_format="pdf")
			payload["volunteer"] = {"name": "Forged Volunteer", "employee": "Forged Employee"}
			payload["vendor_will_sign"] = True
			payload["vendor_signature_data"] = _signature_data()
			payload["authorised_signatory"] = "Asha Vendor"
			with_vendor = invoices.generate_invoice_documents(
				payload, output_format="pdf", generation_reference=volunteer_only["generation_reference"]
			)
			word = invoices.generate_invoice_documents(
				payload, output_format="docx", generation_reference=with_vendor["generation_reference"]
			)
		self.assertEqual(volunteer_only["invoice_number"], "INV-2098-000001")
		self.assertEqual(with_vendor["invoice_number"], "INV-2098-000002")
		self.assertEqual(word["invoice_number"], with_vendor["invoice_number"])
		html = base64.b64decode(with_vendor["pdf"]["content_base64"]).decode()
		self.assertIn("Test Volunteer", html)
		self.assertIn("Test Employee", html)
		self.assertNotIn("Forged", html)
		self.assertIn("not registered", html)
		self.assertTrue(with_vendor["vendor_signed"])
		self.assertEqual(with_vendor["notice"], "")

	def test_changed_content_gets_new_number_not_reference_number(self):
		payload = _payload()
		with self.generation_context():
			first = invoices.generate_invoice_documents(payload, output_format="pdf")
			payload["items"][0]["rate"] = 200
			second = invoices.generate_invoice_documents(
				payload, output_format="docx", generation_reference=first["generation_reference"]
			)
		self.assertEqual(second["invoice_number"], "INV-2098-000002")

	def test_reference_cannot_be_forged_or_used_by_another_employee(self):
		with self.generation_context():
			first = invoices.generate_invoice_documents(_payload(), output_format="pdf")
			with self.assertRaises(frappe.ValidationError):
				invoices.generate_invoice_documents(
					_payload(), output_format="docx", generation_reference="FORGED"
				)
			with patch.object(invoices, "_require_employee", return_value="Another Employee"):
				with self.assertRaises(frappe.PermissionError):
					invoices.generate_invoice_documents(
						_payload(), output_format="docx", generation_reference=first["generation_reference"]
					)

	def test_invalid_format_does_not_consume_number(self):
		with self.generation_context():
			with self.assertRaises(frappe.ValidationError):
				invoices.generate_invoice_documents(_payload(), output_format="invalid")
			self.assertEqual(invoices._next_invoice_number(), "INV-2098-000001")

	def test_number_is_server_assigned_and_identical_in_both_documents_and_filenames(self):
		from docx import Document

		payload = _payload()
		payload["invoice_number"] = "FORGED-NUMBER-THAT-IS-TOO-LONG-FOR-GST"
		with self.generation_context():
			result = invoices.generate_invoice_documents(payload)
		self.assertEqual(result["invoice_number"], "INV-2098-000001")
		html = base64.b64decode(result["pdf"]["content_base64"]).decode()
		word = Document(BytesIO(base64.b64decode(result["docx"]["content_base64"])))
		text = "\n".join(cell.text for table in word.tables for row in table.rows for cell in row.cells)
		for content in (html, text):
			self.assertIn(result["invoice_number"], content)
			self.assertNotIn("FORGED-NUMBER", content)
		self.assertEqual(result["pdf"]["filename"], "invoice-INV-2098-000001.pdf")
		self.assertEqual(result["docx"]["filename"], "invoice-INV-2098-000001.docx")

	def test_invalid_payload_does_not_allocate_a_number(self):
		payload = _payload()
		payload.pop("invoice_number")
		payload["items"] = []
		with self.generation_context():
			with self.assertRaises(frappe.ValidationError):
				invoices.generate_invoice_documents(payload)
			self.assertEqual(
				invoices.generate_invoice_documents(_payload())["invoice_number"], "INV-2098-000001"
			)

	def test_failed_document_generation_rolls_back_number_with_request_transaction(self):
		frappe.db.savepoint("failed_generation_request")
		with (
			self.generation_context(),
			patch.object(invoices, "_build_docx", side_effect=RuntimeError("Render failed")),
		):
			with self.assertRaisesRegex(RuntimeError, "Render failed"):
				invoices.generate_invoice_documents(_payload())
		# Frappe rolls back unsuccessful POST requests, including the Series write.
		frappe.db.rollback(save_point="failed_generation_request")
		with self.generation_context():
			self.assertEqual(
				invoices.generate_invoice_documents(_payload())["invoice_number"], "INV-2098-000001"
			)
