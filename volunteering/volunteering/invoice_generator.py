# Copyright (c) 2026, Vadiraj Tirtha Das and contributors
# For license information, please see license.txt

"""Employee-facing GST and non-GST invoice document generator.

This module deliberately does not create an ERP accounting document. It returns
an in-memory PDF and DOCX carrying the employee's mandatory reimbursement
declaration and, when requested, an additional supplier declaration/signature.
"""

from __future__ import annotations

import base64
import hashlib
import json
import re
import struct
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from io import BytesIO

import frappe
from frappe import _
from frappe.model.naming import getseries
from frappe.utils import cint, formatdate, getdate, money_in_words, now_datetime, nowdate
from frappe.utils.password import decrypt, encrypt
from markupsafe import escape

from volunteering.volunteering.authority import get_employee_for_user

MAX_ITEMS = 20
MAX_SIGNATURE_BYTES = 500_000
VENDOR_ADDRESS_DOCTYPE = "Employee Vendor Address"
VENDOR_STYLE_DOCTYPE = "Vendor Invoice Style"
EMPLOYEE_SIGNATURE_DOCTYPE = "Employee Invoice Signature"
MONEY_PLACES = Decimal("0.01")
MAX_MONEY = Decimal("999999999999.99")
MAX_QUANTITY = Decimal("999999999")
ALLOWED_TYPES = {"GST", "NON_GST"}
ALLOWED_TAX_MODES = {"CGST_SGST", "IGST"}
GSTIN_PATTERN = re.compile(r"^[0-9]{2}[A-Z]{5}[0-9]{4}[A-Z][1-9A-Z]Z[0-9A-Z]$")
PAN_PATTERN = re.compile(r"^[A-Z]{5}[0-9]{4}[A-Z]$")

# Twenty restrained, print-friendly themes. A vendor receives one theme on its
# first use and keeps it thereafter; see _vendor_invoice_style().
INVOICE_STYLES = (
	{
		"id": "style-01",
		"name": "Emerald Ledger",
		"accent": "16745B",
		"soft": "E7F4EF",
		"font": "Arial",
		"align": "center",
		"frame": "line",
	},
	{
		"id": "style-02",
		"name": "Sapphire Header",
		"accent": "245C9A",
		"soft": "EAF1FA",
		"font": "Arial",
		"align": "left",
		"frame": "line",
	},
	{
		"id": "style-03",
		"name": "Maroon Register",
		"accent": "8A3342",
		"soft": "F8ECEE",
		"font": "Georgia",
		"align": "center",
		"frame": "line",
	},
	{
		"id": "style-04",
		"name": "Amber Statement",
		"accent": "A45C09",
		"soft": "FCF2E3",
		"font": "Trebuchet MS",
		"align": "right",
		"frame": "line",
	},
	{
		"id": "style-05",
		"name": "Forest Band",
		"accent": "285C3C",
		"soft": "EAF2ED",
		"font": "Arial",
		"align": "left",
		"frame": "band",
	},
	{
		"id": "style-06",
		"name": "Indigo Band",
		"accent": "4B4E9B",
		"soft": "EEEEF9",
		"font": "Trebuchet MS",
		"align": "center",
		"frame": "band",
	},
	{
		"id": "style-07",
		"name": "Terracotta Band",
		"accent": "A34F32",
		"soft": "F8EDE8",
		"font": "Georgia",
		"align": "left",
		"frame": "band",
	},
	{
		"id": "style-08",
		"name": "Slate Band",
		"accent": "465867",
		"soft": "EDF1F4",
		"font": "Arial",
		"align": "right",
		"frame": "band",
	},
	{
		"id": "style-09",
		"name": "Teal Box",
		"accent": "087D79",
		"soft": "E5F5F4",
		"font": "Trebuchet MS",
		"align": "center",
		"frame": "box",
	},
	{
		"id": "style-10",
		"name": "Royal Box",
		"accent": "5A3E9B",
		"soft": "F0ECFA",
		"font": "Georgia",
		"align": "left",
		"frame": "box",
	},
	{
		"id": "style-11",
		"name": "Copper Box",
		"accent": "97602E",
		"soft": "F6EFE8",
		"font": "Arial",
		"align": "center",
		"frame": "box",
	},
	{
		"id": "style-12",
		"name": "Navy Box",
		"accent": "23466D",
		"soft": "E9EFF5",
		"font": "Trebuchet MS",
		"align": "right",
		"frame": "box",
	},
	{
		"id": "style-13",
		"name": "Olive Double",
		"accent": "65712D",
		"soft": "F1F3E6",
		"font": "Georgia",
		"align": "left",
		"frame": "double",
	},
	{
		"id": "style-14",
		"name": "Plum Double",
		"accent": "7B3F78",
		"soft": "F5EBF4",
		"font": "Arial",
		"align": "center",
		"frame": "double",
	},
	{
		"id": "style-15",
		"name": "Ocean Double",
		"accent": "146B84",
		"soft": "E7F3F6",
		"font": "Trebuchet MS",
		"align": "right",
		"frame": "double",
	},
	{
		"id": "style-16",
		"name": "Brick Double",
		"accent": "884237",
		"soft": "F6ECEA",
		"font": "Georgia",
		"align": "center",
		"frame": "double",
	},
	{
		"id": "style-17",
		"name": "Graphite Minimal",
		"accent": "3F474E",
		"soft": "F0F2F3",
		"font": "Arial",
		"align": "left",
		"frame": "minimal",
	},
	{
		"id": "style-18",
		"name": "Blue Minimal",
		"accent": "286AA6",
		"soft": "EAF2F9",
		"font": "Trebuchet MS",
		"align": "center",
		"frame": "minimal",
	},
	{
		"id": "style-19",
		"name": "Wine Minimal",
		"accent": "883B59",
		"soft": "F7EBF0",
		"font": "Georgia",
		"align": "right",
		"frame": "minimal",
	},
	{
		"id": "style-20",
		"name": "Pine Minimal",
		"accent": "2F6B57",
		"soft": "EAF3EF",
		"font": "Arial",
		"align": "center",
		"frame": "minimal",
	},
)
INVOICE_STYLE_BY_ID = {style["id"]: style for style in INVOICE_STYLES}


@frappe.whitelist(methods=["POST"])
def get_invoice_generator_defaults():
	employee = _require_employee()
	company = frappe.db.get_value("Employee", employee, "company")
	company_doc = frappe.get_cached_doc("Company", company)
	address_choices = _company_address_choices(company, company_doc)
	default_choice = next(
		(choice for choice in address_choices if choice["is_primary_address"]),
		address_choices[0] if address_choices else None,
	)
	address = default_choice["party"] if default_choice else _company_party(company_doc, {})
	return {
		"employee": employee,
		"volunteer": _employee_signer(employee),
		"company": company,
		"invoice_date": nowdate(),
		"office_addresses": address_choices,
		"default_office_address": default_choice["name"] if default_choice else "",
		"consignee": address,
		"vendor_addresses": _vendor_address_choices(employee),
		"saved_volunteer_signature": _saved_employee_signature(employee),
	}


@frappe.whitelist(methods=["POST"])
def generate_invoice_documents(payload, output_format="both", generation_reference=None):
	"""Validate invoice data and return private, in-memory PDF and DOCX downloads."""
	# MariaDB snapshot isolation may reject a counter updated by an overlapping
	# request even after a row lock. Restart the whole generation transaction and
	# re-check employee access rather than retrying with a stale snapshot.
	for attempt in range(3):
		try:
			return _generate_invoice_documents(payload, output_format, generation_reference)
		except frappe.QueryDeadlockError:
			frappe.db.rollback()
			if attempt == 2:
				frappe.throw(
					_("Invoice generation is busy. Please try again; no invoice number was consumed.")
				)


def _generate_invoice_documents(payload, output_format="both", generation_reference=None):
	if output_format not in {"pdf", "docx", "both"}:
		frappe.throw(_("Choose PDF or Word document."))
	employee = _require_employee()
	# Validate first, without trusting or allocating a browser-supplied number.
	data = _normalise_payload(
		_apply_selected_office_addresses(payload, employee),
		invoice_number_override="INV-AUTOMATIC",
		volunteer_override=_employee_signer(employee),
	)
	vendor_identity_key, invoice_style = _vendor_invoice_style(employee, data["supplier"])
	data["invoice_style"] = invoice_style
	fingerprint_data = {
		key: value
		for key, value in data.items()
		if key not in {"volunteer_signature_png", "vendor_signature_png"}
	}
	fingerprint_data["volunteer_signature_sha256"] = hashlib.sha256(
		data["volunteer_signature_png"]
	).hexdigest()
	fingerprint_data["vendor_signature_sha256"] = (
		hashlib.sha256(data["vendor_signature_png"]).hexdigest()
		if data["vendor_signature_png"]
		else ""
	)
	fingerprint = hashlib.sha256(
		json.dumps(fingerprint_data, sort_keys=True, default=str, separators=(",", ":")).encode()
	).hexdigest()
	data["invoice_number"] = _number_for_generation(employee, fingerprint, generation_reference)
	stem = _safe_filename(data["invoice_number"])
	result = {
		"invoice_number": data["invoice_number"],
		"supplier_name": data["supplier"]["name"],
		"grand_total": float(data["grand_total"]),
		"amount_in_words": data["amount_in_words"],
		"vendor_signed": data["vendor_will_sign"],
		"notice": data["notice"],
		"invoice_style": data["invoice_style"],
	}
	if output_format in {"pdf", "both"}:
		result["pdf"] = _download(f"{stem}.pdf", "application/pdf", _build_pdf(data))
	if output_format in {"docx", "both"}:
		result["docx"] = _download(
			f"{stem}.docx",
			"application/vnd.openxmlformats-officedocument.wordprocessingml.document",
			_build_docx(data),
		)
	# A successfully rendered invoice is the point at which a fresh volunteer
	# signature becomes reusable. Failed generation never changes the saved copy.
	_save_employee_signature(employee, data["volunteer_signature_png"])
	# Remember only a successfully rendered invoice. PDF and Word generated from
	# the same opaque reference share one invoice number and count as one use.
	result["vendor_address"] = _remember_vendor_address(
		employee,
		data["supplier"],
		data["invoice_number"],
		bank=data["bank"],
		invoice_style=data["invoice_style"],
		vendor_identity_key=vendor_identity_key,
	)
	# An authenticated, opaque reference permits the other format for exactly
	# this employee and content, without trusting client-supplied invoice numbers.
	# Issue it only after the requested document(s) succeed.
	result["generation_reference"] = encrypt(
		json.dumps(
			{
				"purpose": "invoice-generator-v1",
				"employee": employee,
				"fingerprint": fingerprint,
				"invoice_number": data["invoice_number"],
			}
		)
	)
	return result


def _vendor_address_key(employee, party):
	identity = "\n".join(
		[
			_text(employee, 140).casefold(),
			_text(party.get("name"), 160).casefold(),
			_text(party.get("address"), 600).casefold(),
			_text(party.get("state"), 100).casefold(),
			_text(party.get("pin_code"), 12).casefold(),
		]
	)
	return hashlib.sha256(identity.encode()).hexdigest()


def _normalised_vendor_identity(value):
	return re.sub(r"[^a-z0-9]+", " ", str(value or "").casefold()).strip()


def _vendor_identity_key(party):
	# The legal name is the stable identity employees select again. Addresses,
	# registrations and remittance details may legitimately change without making
	# the supplier a different visual identity.
	identity = f"name:{_normalised_vendor_identity(party.get('name'))}"
	return hashlib.sha256(identity.encode()).hexdigest()


def _choose_invoice_style(assignments):
	"""Choose the least-used style; ties preserve the published 1-20 order."""
	counts = {style["id"]: 0 for style in INVOICE_STYLES}
	seen = set()
	for row in assignments:
		key = _normalised_vendor_identity(row.get("vendor_name")) or row.get("vendor_key") or row.get("name")
		style = row.get("invoice_style")
		if key and key not in seen and style in counts:
			seen.add(key)
			counts[style] += 1
	return min(INVOICE_STYLES, key=lambda style: (counts[style["id"]], style["id"]))["id"]


def _vendor_invoice_style(employee, party):
	"""Return the organisation-wide stable style assigned to this vendor identity."""
	vendor_key = _vendor_identity_key(party)
	style = frappe.db.get_value(VENDOR_STYLE_DOCTYPE, vendor_key, "invoice_style")
	if style in INVOICE_STYLE_BY_ID:
		return vendor_key, style
	# Preserve the earliest assignment created by the first implementation, whose
	# key could include an address or GSTIN. This also collapses such legacy
	# duplicates to one effective vendor for future generation.
	vendor_name = _normalised_vendor_identity(party.get("name"))
	legacy_assignments = frappe.get_all(
		VENDOR_STYLE_DOCTYPE,
		fields=["name", "vendor_name", "invoice_style"],
		order_by="assignment_sequence asc, creation asc",
		limit_page_length=0,
	)
	for assignment in legacy_assignments:
		if (
			_normalised_vendor_identity(assignment.vendor_name) == vendor_name
			and assignment.invoice_style in INVOICE_STYLE_BY_ID
		):
			return assignment.name, assignment.invoice_style

	company = frappe.db.get_value("Employee", employee, "company")
	if company:
		# Serialise first-time assignments so two simultaneous new vendors cannot
		# consume the same position before all twenty styles have been used.
		frappe.db.sql("SELECT name FROM `tabCompany` WHERE name=%s FOR UPDATE", (company,))
	style = frappe.db.get_value(VENDOR_STYLE_DOCTYPE, vendor_key, "invoice_style")
	if style not in INVOICE_STYLE_BY_ID:
		assignments = frappe.get_all(
			VENDOR_STYLE_DOCTYPE,
			fields=["vendor_key", "vendor_name", "invoice_style"],
			order_by="assignment_sequence asc, creation asc",
			limit_page_length=0,
		)
		style = _choose_invoice_style(assignments)
		frappe.get_doc(
			{
				"doctype": VENDOR_STYLE_DOCTYPE,
				"vendor_key": vendor_key,
				"vendor_name": party["name"],
				"vendor_gstin": party.get("gstin") or "",
				"invoice_style": style,
				"assignment_sequence": len({row.vendor_key for row in assignments}) + 1,
			}
		).insert(ignore_permissions=True)
	return vendor_key, style


def _vendor_address_row(row):
	party = {
		"name": row.vendor_name,
		"address": row.address,
		"state": row.state,
		"pin_code": row.pin_code or "",
		"gstin": row.gstin or "",
		"pan": row.pan or "",
	}
	location = ", ".join(part for part in (party["address"], party["state"], party["pin_code"]) if part)
	account_number = row.get_password("account_number", raise_exception=False) or ""
	bank = {
		"bank_name": row.bank_name or "",
		"branch": row.bank_branch or "",
		"account_name": row.account_name or "",
		"account_number": account_number,
		"ifsc": row.ifsc or "",
		"swift": row.swift or "",
		"upi_id": row.upi_id or "",
	}
	return {
		"name": row.name,
		"label": f"{party['name']} — {location}",
		"party": party,
		"bank": bank,
		"invoice_style": row.invoice_style or "",
		"use_count": row.use_count or 0,
		"last_used_on": row.last_used_on,
	}


def _vendor_address_choices(employee):
	names = frappe.get_all(
		VENDOR_ADDRESS_DOCTYPE,
		filters={"employee": employee},
		pluck="name",
		order_by="use_count desc, last_used_on desc, modified desc",
		limit=100,
	)
	rows = []
	for name in names:
		doc = frappe.get_doc(VENDOR_ADDRESS_DOCTYPE, name)
		vendor_key, style = _vendor_invoice_style(
			employee,
			{
				"name": doc.vendor_name,
				"address": doc.address,
				"state": doc.state,
				"pin_code": doc.pin_code,
				"gstin": doc.gstin,
			},
		)
		if doc.vendor_identity_key != vendor_key or doc.invoice_style != style:
			doc.vendor_identity_key = vendor_key
			doc.invoice_style = style
			doc.save(ignore_permissions=True)
		rows.append(_vendor_address_row(doc))
	return rows


def _remember_vendor_address(
	employee,
	party,
	invoice_number,
	bank=None,
	invoice_style=None,
	vendor_identity_key=None,
):
	key = _vendor_address_key(employee, party)
	bank = _bank(bank)
	if not vendor_identity_key or invoice_style not in INVOICE_STYLE_BY_ID:
		vendor_identity_key, invoice_style = _vendor_invoice_style(employee, party)
	doc = frappe.db.exists(VENDOR_ADDRESS_DOCTYPE, key)
	if doc:
		doc = frappe.get_doc(VENDOR_ADDRESS_DOCTYPE, doc)
		if doc.employee != employee:
			frappe.throw(_("The saved vendor address belongs to another employee."), frappe.PermissionError)
		doc.vendor_name = party["name"]
		doc.address = party["address"]
		doc.state = party["state"]
		doc.pin_code = party.get("pin_code") or ""
		doc.gstin = party.get("gstin") or ""
		doc.pan = party.get("pan") or ""
		doc.vendor_identity_key = vendor_identity_key
		doc.invoice_style = invoice_style
		doc.bank_name = bank["bank_name"]
		doc.bank_branch = bank["branch"]
		doc.account_name = bank["account_name"]
		doc.account_number = bank["account_number"]
		doc.ifsc = bank["ifsc"]
		doc.swift = bank["swift"]
		doc.upi_id = bank["upi_id"]
		if doc.last_invoice_number != invoice_number:
			doc.use_count = (doc.use_count or 0) + 1
	else:
		doc = frappe.get_doc(
			{
				"doctype": VENDOR_ADDRESS_DOCTYPE,
				"address_key": key,
				"employee": employee,
				"vendor_name": party["name"],
				"address": party["address"],
				"state": party["state"],
				"pin_code": party.get("pin_code") or "",
				"gstin": party.get("gstin") or "",
				"pan": party.get("pan") or "",
				"vendor_identity_key": vendor_identity_key,
				"invoice_style": invoice_style,
				"bank_name": bank["bank_name"],
				"bank_branch": bank["branch"],
				"account_name": bank["account_name"],
				"account_number": bank["account_number"],
				"ifsc": bank["ifsc"],
				"swift": bank["swift"],
				"upi_id": bank["upi_id"],
				"use_count": 1,
			}
		)
	doc.last_invoice_number = invoice_number
	doc.last_used_on = now_datetime()
	if doc.is_new():
		doc.insert(ignore_permissions=True)
	else:
		doc.save(ignore_permissions=True)
	return _vendor_address_row(doc)


def _number_for_generation(employee, fingerprint, reference):
	if reference:
		try:
			identity = json.loads(decrypt(reference))
		except ValueError, TypeError, frappe.ValidationError:
			frappe.throw(_("Invalid invoice generation reference. Generate again."))
		if not isinstance(identity, dict) or identity.get("purpose") != "invoice-generator-v1":
			frappe.throw(_("Invalid invoice generation reference. Generate again."))
		if identity.get("employee") != employee:
			frappe.throw(_("This invoice belongs to another employee."), frappe.PermissionError)
		if identity.get("fingerprint") == fingerprint:
			return identity["invoice_number"]
	return _next_invoice_number()


def _next_invoice_number():
	"""Allocate one shared annual number inside the generation transaction."""
	year = getdate(nowdate()).year
	key = f"SEVAMRITA-FORM-INVOICE-{year}-"
	series = frappe.qb.DocType("Series")
	if not frappe.qb.from_(series).select(series.name).where(series.name == key).run():
		# The first generation has no Series row to lock. Serialize its creation
		# on an existing site Company row, then let getseries re-read with a lock.
		# Avoid a no-op UPSERT: MariaDB can reject it from a stale snapshot.
		company = frappe.qb.DocType("Company")
		frappe.qb.from_(company).select(company.name).orderby(company.name).limit(1).for_update().run()
	# getseries locks the counter until both documents succeed and the POST
	# commits. Frappe rolls back allocation if document generation fails.
	number = f"INV-{year}-{getseries(key, 6)}"
	if len(number) > 16:
		frappe.throw(_("The annual invoice numbering sequence is exhausted. Contact your administrator."))
	return number


def _require_employee():
	user = frappe.session.user
	if not user or user == "Guest":
		frappe.throw(_("Log in as an employee to prepare an invoice."), frappe.PermissionError)
	employee = get_employee_for_user(user)
	if not employee or frappe.db.get_value("Employee", employee, "status") != "Active":
		frappe.throw(_("Your user must be linked to an active Employee record."), frappe.PermissionError)
	return employee


def _employee_signer(employee):
	# A volunteer signs their own statement. Never trust a browser-supplied
	# employee identity or present the volunteer as a supplier representative.
	return {
		"employee": employee,
		"name": frappe.db.get_value("Employee", employee, "employee_name") or employee,
	}


def _saved_employee_signature(employee):
	"""Return only the current employee's reusable signature as a data URL."""
	if not frappe.db.exists("DocType", EMPLOYEE_SIGNATURE_DOCTYPE):
		return ""
	encoded = frappe.db.get_value(EMPLOYEE_SIGNATURE_DOCTYPE, employee, "signature_data") or ""
	if not encoded:
		return ""
	value = f"data:image/png;base64,{encoded}"
	try:
		_signature_png(value)
	except frappe.ValidationError:
		return ""
	return value


def _save_employee_signature(employee, signature_png):
	"""Privately retain the last valid volunteer signature for explicit reuse."""
	if not signature_png:
		return
	encoded = base64.b64encode(signature_png).decode("ascii")
	doc = frappe.db.exists(EMPLOYEE_SIGNATURE_DOCTYPE, employee)
	if doc:
		doc = frappe.get_doc(EMPLOYEE_SIGNATURE_DOCTYPE, doc)
		doc.signature_data = encoded
		doc.signature_sha256 = hashlib.sha256(signature_png).hexdigest()
		doc.last_used_on = now_datetime()
		doc.save(ignore_permissions=True)
	else:
		frappe.get_doc(
			{
				"doctype": EMPLOYEE_SIGNATURE_DOCTYPE,
				"employee": employee,
				"signature_data": encoded,
				"signature_sha256": hashlib.sha256(signature_png).hexdigest(),
				"last_used_on": now_datetime(),
			}
		).insert(ignore_permissions=True)


def _company_address_choices(company, company_doc=None):
	from volunteering.volunteering.office_addresses import invoice_address_choices

	company_doc = company_doc or frappe.get_cached_doc("Company", company)
	choices = invoice_address_choices(company)
	for choice in choices:
		choice["party"] = _company_party(company_doc, choice["party"])
	return choices


def _company_party(company_doc, address):
	return {
		"name": company_doc.company_name or company_doc.name,
		"address": address.get("address", ""),
		"state": address.get("state", ""),
		"pin_code": address.get("pin_code", ""),
		"gstin": (company_doc.get("tax_id") or "").strip().upper(),
	}


def _apply_selected_office_addresses(payload, employee):
	raw = frappe.parse_json(payload) if isinstance(payload, str) else payload
	if not isinstance(raw, dict):
		return raw
	consignee_name = str(raw.get("consignee_address_name") or "").strip()
	if not consignee_name:
		return raw

	company = frappe.db.get_value("Employee", employee, "company")
	choices = {choice["name"]: choice for choice in _company_address_choices(company)}
	values = dict(raw)
	if consignee_name:
		if consignee_name not in choices:
			frappe.throw(_("Choose a current Sevamrita office address for the consignee."))
		values["consignee"] = choices[consignee_name]["party"]
	# Buyer and consignee are one Sevamrita office for these employee-generated
	# invoices. Browser-supplied alternate buyer data is deliberately ignored.
	values["buyer"] = values["consignee"]
	values["buyer_same_as_consignee"] = True
	return values


def _normalise_payload(payload, bank_override=None, invoice_number_override=None, volunteer_override=None):
	raw = frappe.parse_json(payload) if isinstance(payload, str) else payload
	if not isinstance(raw, dict):
		frappe.throw(_("Invoice details must be a valid object."))

	invoice_type = _required_choice(raw, "invoice_type", ALLOWED_TYPES, _("Invoice type"))
	volunteer = volunteer_override if volunteer_override is not None else raw.get("volunteer")
	volunteer = volunteer if isinstance(volunteer, dict) else {}
	vendor_will_sign = bool(cint(raw.get("vendor_will_sign")))
	# During a rolling deployment, accept the former VOLUNTEER signature field as
	# the mandatory volunteer signature. Supplier-only legacy payloads remain
	# invalid because the reimbursement declaration must now always be signed.
	legacy_volunteer_signature = (
		raw.get("signature_data") if str(raw.get("signer_type") or "").upper() == "VOLUNTEER" else None
	)
	volunteer_signature_png = _signature_png(
		raw.get("volunteer_signature_data") or legacy_volunteer_signature
	)
	if not volunteer_signature_png:
		frappe.throw(_("The volunteer signature is required for reimbursement."))
	vendor_signature_png = _signature_png(raw.get("vendor_signature_data")) if vendor_will_sign else b""
	if vendor_will_sign and not vendor_signature_png:
		frappe.throw(_("Ask the vendor to sign on screen, or turn off vendor signing."))
	data = {
		"invoice_type": invoice_type,
		"vendor_will_sign": vendor_will_sign,
		"copy_label": _("ORIGINAL FOR RECIPIENT"),
		"title": "TAX INVOICE" if invoice_type == "GST" else "INVOICE",
		"invoice_number": invoice_number_override
		if invoice_number_override is not None
		else _required_text(raw, "invoice_number", _("Invoice number"), 80),
		"invoice_date": _date(raw.get("invoice_date"), _("Invoice date")),
		"buyer_order_number": _text(raw.get("buyer_order_number"), 100),
		"buyer_order_date": _optional_date(raw.get("buyer_order_date")),
		"supplier_reference": _text(raw.get("supplier_reference"), 140),
		"dispatch_document_number": _text(raw.get("dispatch_document_number"), 100),
		"delivery_note_date": _optional_date(raw.get("delivery_note_date")),
		"supplier": _party(raw.get("supplier"), _("Supplier")),
		"consignee": _party(raw.get("consignee"), _("Consignee")),
		"buyer_same_as_consignee": True,
		"transportation_charges": _money(raw.get("transportation_charges"), _("Transportation charges")),
		"other_charges": _money(raw.get("other_charges"), _("Other charges")),
		# These are optional supplier remittance details. Employee reimbursement
		# banking is deliberately separate and is never printed on a supplier invoice.
		"bank": _bank(bank_override if bank_override is not None else raw.get("bank")),
		"invoice_style": "style-01",
		"authorised_signatory": _text(raw.get("authorised_signatory"), 120)
		if vendor_will_sign
		else "",
		"volunteer_signature_png": volunteer_signature_png,
		"vendor_signature_png": vendor_signature_png,
	}
	data["buyer"] = data["consignee"]

	if invoice_type == "GST":
		_validate_gstin(data["supplier"].get("gstin"), _("Supplier GSTIN"), required=True)
		_validate_gstin(data["consignee"].get("gstin"), _("Consignee GSTIN"), required=False)
		if len(data["invoice_number"]) > 16:
			frappe.throw(_("A GST invoice number cannot exceed 16 characters."))
		# The supplied Tax Invoice template has one GST amount row, not separate
		# tax-mode/rate/place-of-supply fields. Accept the amount shown on the bill.
		if "gst_amount" in raw:
			data["gst_amount"] = _money(raw.get("gst_amount"), _("GST amount"))
		else:
			# Compatibility for saved/older clients while the home form migrates.
			legacy_rate = _rate(raw.get("gst_rate"))
			data["gst_amount"] = None
		data["tax_mode"] = _text(raw.get("tax_mode"), 20)
		data["gst_rate"] = _rate(raw.get("gst_rate")) if raw.get("gst_rate") is not None else Decimal("0")
	else:
		_validate_pan(data["supplier"].get("pan"))
		data["tax_mode"] = ""
		data["gst_rate"] = Decimal("0")
		data["gst_amount"] = Decimal("0")

	items = raw.get("items")
	if not isinstance(items, list) or not items:
		frappe.throw(_("Add at least one invoice item."))
	if len(items) > MAX_ITEMS:
		frappe.throw(_("An invoice can contain at most {0} items.").format(MAX_ITEMS))

	data["items"] = []
	for index, raw_item in enumerate(items, start=1):
		if not isinstance(raw_item, dict):
			frappe.throw(_("Invoice item {0} is invalid.").format(index))
		quantity = _positive_decimal(raw_item.get("quantity"), _("Quantity in row {0}").format(index), 3)
		rate = _money(raw_item.get("rate"), _("Rate in row {0}").format(index))
		if rate <= 0:
			frappe.throw(_("Rate in row {0} must be greater than zero.").format(index))
		line_amount = quantity * rate
		if line_amount > MAX_MONEY:
			frappe.throw(_("Amount in row {0} is too large.").format(index))
		data["items"].append(
			{
				"number": index,
				"description": _required_text(
					raw_item, "description", _("Description in row {0}").format(index), 300
				),
				"hsn_sac": _text(raw_item.get("hsn_sac"), 30),
				"unit": _required_text(raw_item, "unit", _("Unit in row {0}").format(index), 20),
				"quantity": quantity,
				"rate": rate,
				"amount": _quantise(line_amount),
			}
		)

	data["items_total"] = _quantise(sum((item["amount"] for item in data["items"]), Decimal("0")))
	taxable_total = data["items_total"] + data["transportation_charges"] + data["other_charges"]
	if taxable_total > MAX_MONEY:
		frappe.throw(_("Invoice total is too large."))
	data["taxable_total"] = _quantise(taxable_total)
	if data["gst_amount"] is None:
		data["gst_amount"] = _quantise(data["taxable_total"] * legacy_rate / Decimal("100"))
	else:
		data["gst_amount"] = _quantise(data["gst_amount"])
	data["cgst_amount"] = Decimal("0")
	data["sgst_amount"] = Decimal("0")
	data["igst_amount"] = Decimal("0")
	if data["tax_mode"] == "CGST_SGST":
		data["cgst_amount"] = _quantise(data["gst_amount"] / Decimal("2"))
		data["sgst_amount"] = data["gst_amount"] - data["cgst_amount"]
	elif data["tax_mode"] == "IGST":
		data["igst_amount"] = data["gst_amount"]
	grand_total = data["taxable_total"] + data["gst_amount"]
	if grand_total > MAX_MONEY:
		frappe.throw(_("Invoice total is too large."))
	data["grand_total"] = _quantise(grand_total)
	data["amount_in_words"] = money_in_words(data["grand_total"], "INR")
	data["volunteer_name"] = _required_text(volunteer, "name", _("Volunteer name"), 160)
	data["volunteer_employee"] = _required_text(volunteer, "employee", _("Volunteer employee"), 140)
	data["volunteer_declaration"] = _(
		"I confirm that I paid the amount shown above and request reimbursement to my bank account."
	)
	data["vendor_declaration"] = ""
	if vendor_will_sign and invoice_type == "NON_GST":
		data["vendor_declaration"] = _(
			"I, {0}, proprietor/authorised person of {1}, declare that this business is not registered "
			"under the Goods and Services Tax (GST) Act and therefore does not have a GSTIN."
		).format(data["authorised_signatory"] or _("the undersigned"), data["supplier"]["name"])
	data["notice"] = ""
	return data


def _signature_png(value):
	"""Accept only a modest, non-blank PNG created by the signature canvas."""
	if not value:
		return b""
	if not isinstance(value, str) or not value.startswith("data:image/png;base64,"):
		frappe.throw(_("The on-screen signature must be a PNG image."))
	encoded = value.partition(",")[2]
	if len(encoded) > ((MAX_SIGNATURE_BYTES * 4) // 3) + 8:
		frappe.throw(_("The on-screen signature is too large. Clear it and sign again."))
	try:
		content = base64.b64decode(encoded, validate=True)
	except Exception:
		frappe.throw(_("The on-screen signature image is invalid."))
	if len(content) > MAX_SIGNATURE_BYTES or not content.startswith(b"\x89PNG\r\n\x1a\n"):
		frappe.throw(_("The on-screen signature image is invalid."))
	if len(content) < 24:
		frappe.throw(_("The on-screen signature image is invalid."))
	width, height = struct.unpack(">II", content[16:24])
	if width < 120 or height < 40 or width > 1600 or height > 800 or width * height > 1_280_000:
		frappe.throw(_("The on-screen signature has invalid dimensions."))
	try:
		from PIL import Image, ImageChops

		with Image.open(BytesIO(content)) as image:
			image.load()
			if image.format != "PNG" or image.size != (width, height):
				raise ValueError
			# Composite transparency over the canvas's white background so a fully
			# transparent PNG cannot masquerade as a non-blank signature.
			rgba = image.convert("RGBA")
			white = Image.new("RGBA", rgba.size, "white")
			white.alpha_composite(rgba)
			rgb = white.convert("RGB")
			if ImageChops.difference(rgb, Image.new("RGB", rgb.size, "white")).getbbox() is None:
				frappe.throw(_("Draw a signature before saving it."))
	except frappe.ValidationError:
		raise
	except Exception:
		frappe.throw(_("The on-screen signature image is invalid."))
	return content


def _party(value, label):
	value = value if isinstance(value, dict) else {}
	return {
		"name": _required_text(value, "name", _("{0} name").format(label), 160),
		"address": _required_text(value, "address", _("{0} address").format(label), 600),
		"state": _required_text(value, "state", _("{0} state").format(label), 100),
		"pin_code": _text(value.get("pin_code"), 12),
		"gstin": _text(value.get("gstin"), 15).upper(),
		"pan": _text(value.get("pan"), 10).upper(),
	}


def _bank(value):
	value = value if isinstance(value, dict) else {}
	return {
		"bank_name": _text(value.get("bank_name"), 140),
		"branch": _text(value.get("branch"), 140),
		"account_name": _text(value.get("account_name"), 160),
		"account_number": _text(value.get("account_number"), 50),
		"ifsc": _text(value.get("ifsc"), 20).upper(),
		"swift": _text(value.get("swift"), 20).upper(),
		"upi_id": _text(value.get("upi_id"), 100),
	}


def _has_bank(bank):
	return any(
		bank.get(key)
		for key in ("bank_name", "branch", "account_name", "account_number", "ifsc", "swift", "upi_id")
	)


def _bank_lines(bank):
	return [
		(_("Bank Name"), bank["bank_name"]),
		(_("Branch"), bank["branch"]),
		(_("Account Name"), bank["account_name"]),
		(_("Account No."), bank["account_number"]),
		(_("IFSC Code"), bank["ifsc"]),
		(_("Swift Code"), bank["swift"]),
		(_("UPI ID"), bank["upi_id"]),
	]


def _invoice_style(data):
	return INVOICE_STYLE_BY_ID.get(data.get("invoice_style"), INVOICE_STYLES[0])


def _required_text(mapping, key, label, max_length):
	value = _text(mapping.get(key), max_length)
	if not value:
		frappe.throw(_("{0} is required.").format(label))
	return value


def _text(value, max_length):
	value = " ".join(str(value or "").strip().split())
	if len(value) > max_length:
		frappe.throw(_("A field exceeds the maximum length of {0} characters.").format(max_length))
	return value


def _required_choice(mapping, key, choices, label):
	value = _text(mapping.get(key), 40).upper()
	if value not in choices:
		frappe.throw(_("Select a valid {0}.").format(label))
	return value


def _date(value, label):
	if not value:
		frappe.throw(_("{0} is required.").format(label))
	try:
		return formatdate(getdate(value), "dd-MM-yyyy")
	except Exception:
		frappe.throw(_("{0} is invalid.").format(label))


def _optional_date(value):
	return _date(value, _("Date")) if value else ""


def _decimal(value, label):
	try:
		amount = Decimal(str(value or 0))
	except InvalidOperation, ValueError:
		frappe.throw(_("{0} must be a number.").format(label))
	if not amount.is_finite():
		frappe.throw(_("{0} must be a finite number.").format(label))
	return amount


def _money(value, label):
	amount = _decimal(value, label)
	if amount < 0:
		frappe.throw(_("{0} cannot be negative.").format(label))
	if amount > MAX_MONEY:
		frappe.throw(_("{0} is too large.").format(label))
	return _quantise(amount)


def _positive_decimal(value, label, places):
	amount = _decimal(value, label)
	if amount <= 0:
		frappe.throw(_("{0} must be greater than zero.").format(label))
	if amount > MAX_QUANTITY:
		frappe.throw(_("{0} is too large.").format(label))
	quantum = Decimal("1").scaleb(-places)
	return amount.quantize(quantum, rounding=ROUND_HALF_UP)


def _rate(value):
	rate = _decimal(value, _("GST rate"))
	if rate < 0 or rate > 100:
		frappe.throw(_("GST rate must be between 0 and 100."))
	return rate.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)


def _quantise(value):
	return Decimal(value).quantize(MONEY_PLACES, rounding=ROUND_HALF_UP)


def _validate_gstin(value, label, required):
	if not value and not required:
		return
	if not value or not GSTIN_PATTERN.fullmatch(value):
		frappe.throw(_("{0} must be a valid 15-character GSTIN.").format(label))


def _validate_pan(value):
	if value and not PAN_PATTERN.fullmatch(value):
		frappe.throw(_("Supplier PAN must be in the format ABCDE1234F."))


def _money_text(value):
	return f"{Decimal(value):,.2f}"


def _quantity_text(value):
	return f"{Decimal(value):f}".rstrip("0").rstrip(".")


def _download(filename, content_type, content):
	return {
		"filename": filename,
		"content_type": content_type,
		"content_base64": base64.b64encode(content).decode("ascii"),
	}


def _safe_filename(invoice_number):
	stem = re.sub(r"[^A-Za-z0-9._-]+", "-", invoice_number).strip("-._")
	return f"invoice-{stem or 'document'}"


def _render_pdf_html(data):
	def h(value):
		return str(escape(str(value or "")))

	theme = _invoice_style(data)
	title_frame = {
		"line": f"border-bottom:3px solid #{theme['accent']};padding-bottom:3mm",
		"band": f"background:#{theme['accent']};color:#fff;padding:4mm",
		"box": f"border:2px solid #{theme['accent']};padding:3mm",
		"double": f"border-top:4px double #{theme['accent']};border-bottom:4px double #{theme['accent']};padding:3mm",
		"minimal": f"color:#{theme['accent']};letter-spacing:1.5px;padding-bottom:2mm",
	}[theme["frame"]]
	header_fill = theme["accent"] if theme["frame"] == "band" else theme["soft"]
	header_text = "#fff" if theme["frame"] == "band" else "#111"

	items = "".join(
		f"<tr><td>{item['number']}</td><td>{h(item['description'])}</td>"
		f"<td>{h(item['hsn_sac'])}</td><td class='num'>{_quantity_text(item['quantity'])}</td>"
		f"<td class='num'>{_money_text(item['rate'])}</td><td class='num'>{_money_text(item['amount'])}</td></tr>"
		for item in data["items"]
	)
	summary = [
		(_("Total"), data["items_total"]),
	]
	if data["invoice_type"] == "GST":
		summary.append((_("GST"), data["gst_amount"]))
	summary.extend(
		[
			(_("Transportation"), data["transportation_charges"]),
			(_("Others"), data["other_charges"]),
		]
	)
	summary_rows = "".join(
		f"<tr><td colspan='5' class='summary-label'>{h(label)}</td><td class='num'>{_money_text(value)}</td></tr>"
		for label, value in summary
	)
	bank = data["bank"]
	bank_html = "<br>".join(f"<b>{h(label)}:</b> {h(value)}" for label, value in _bank_lines(bank) if value)
	def signature_image(content):
		return (
			f'<img class="signature-image" src="data:image/png;base64,{base64.b64encode(content).decode("ascii")}">'
			if content
			else ""
		)

	bank_section = (
		f'<table class="section"><tr><td><b>{h(_("Supplier Remittance Details"))}'
		f"</b><br>{bank_html}</td></tr></table>"
		if bank_html
		else ""
	)
	volunteer_section = (
		f'<div class="signature-section"><b>{h(_("Volunteer reimbursement declaration"))}</b>'
		f'<p>{h(data["volunteer_declaration"])}</p><div class="signature">'
		f'{signature_image(data["volunteer_signature_png"])}{h(data["volunteer_name"])}'
		f'<br>{h(data["volunteer_employee"])}<br>{h(_("Volunteer Signature"))}</div></div>'
	)
	vendor_section = ""
	if data["vendor_will_sign"]:
		vendor_statement = (
			f'<p>{h(data["vendor_declaration"])}</p>' if data["vendor_declaration"] else ""
		)
		vendor_section = (
			f'<div class="signature-section"><b>{h(_("Vendor confirmation"))}</b>{vendor_statement}'
			f'<div class="signature">{signature_image(data["vendor_signature_png"])}'
			f'{h(data["authorised_signatory"] or data["supplier"]["name"])}'
			f'<br>{h(_("Authorised Signatory for {0}")).format(data["supplier"]["name"])}</div></div>'
		)
	footer_html = bank_section + volunteer_section + vendor_section
	buyer = data["buyer"]
	return f"""<!doctype html>
<html><head><meta charset="utf-8"><style>
@page {{ size: A4; margin: 8mm; }}
body {{ font-family: '{theme["font"]}', sans-serif; color: #111; font-size: 9.5pt; line-height: 1.3; }}
h1 {{ text-align: {theme["align"]}; font-size: 17pt; margin: 0 0 3mm; {title_frame}; }}
.copy {{ text-align: right; font-size: 8pt; margin-bottom: 2mm; }}
table {{ width: 100%; border-collapse: collapse; }}
td, th {{ border: 1px solid #{theme["accent"]}; padding: 4px 5px; vertical-align: top; }}
th {{ background: #{header_fill}; color: {header_text}; text-align: center; }}
thead {{ display: table-header-group; }}
tr {{ page-break-inside: avoid; }}
.items {{ table-layout: fixed; }}
.items td {{ overflow-wrap: break-word; }}
.party {{ width: 50%; }}
.label {{ color: #{theme["accent"]}; font-size: 8pt; text-transform: uppercase; }}
.value {{ font-weight: 600; margin-top: 1px; }}
.num {{ text-align: right; white-space: nowrap; }}
.summary-label {{ text-align: right; font-weight: 600; }}
.grand td {{ font-size: 11pt; font-weight: 700; background: #{theme["soft"]}; }}
.section {{ margin-top: 3mm; }}
.declaration {{ border: 1px solid #{theme["accent"]}; border-left-width: 4px; padding: 5px; margin-top: 3mm; }}
.signature-section {{ border: 1px solid #{theme["accent"]}; border-left-width: 4px; padding: 5px; margin-top: 3mm; page-break-inside: avoid; }}
.signature-section p {{ margin: 2mm 0; }}
.signature {{ min-height: 22mm; text-align: right; }}
.signature-only {{ width: 100%; }}
.signature-image {{ display: block; max-width: 58mm; max-height: 24mm; margin: 1mm 0 1mm auto; object-fit: contain; }}
</style></head><body>
<h1>{h(data["title"])}</h1><div class="copy">{h(data["copy_label"])}</div>
<table><tr><td class="party"><div class="label">{h(_("Supplier Details"))}</div><div class="value">{h(data["supplier"]["name"])}</div>{h(data["supplier"]["address"])}<br>{h(data["supplier"]["state"])} {h(data["supplier"]["pin_code"])}<br>{h(_("GSTIN / UID No.")) if data["invoice_type"] == "GST" else h(_("PAN No."))}: {h(data["supplier"]["gstin"] if data["invoice_type"] == "GST" else data["supplier"]["pan"])}</td>
<td><b>{h(_("Invoice No."))}:</b> {h(data["invoice_number"])}<br><b>{h(_("Invoice Date"))}:</b> {h(data["invoice_date"])}<br><b>{h(_("Buyer's Order No."))}:</b> {h(data["buyer_order_number"])}<br><b>{h(_("Date"))}:</b> {h(data["buyer_order_date"])}<br><b>{h(_("Supplier's Reference"))}:</b> {h(data["supplier_reference"])}<br><b>{h(_("Despatch Document No."))}:</b> {h(data["dispatch_document_number"])}<br><b>{h(_("Delivery Note Date"))}:</b> {h(data["delivery_note_date"])}</td></tr>
<tr><td><div class="label">{h(_("Consignee's Details"))}</div><div class="value">{h(data["consignee"]["name"])}</div>{h(data["consignee"]["address"])}<br>{h(data["consignee"]["state"])} {h(data["consignee"]["pin_code"])}<br>{h(_("GSTIN / UID No."))}: {h(data["consignee"]["gstin"])}</td>
<td><div class="label">{h(_("Buyer's Details (if other than Consignee)"))}</div><div class="value">{h(buyer["name"])}</div>{h(buyer["address"])}<br>{h(buyer["state"])} {h(buyer["pin_code"])}<br>{h(_("GSTIN / UID No."))}: {h(buyer["gstin"])}</td></tr></table>
<table class="section items"><colgroup><col style="width:6%"><col style="width:42%"><col style="width:12%"><col style="width:8%"><col style="width:14%"><col style="width:18%"></colgroup><thead><tr><th>{h(_("Sr. No."))}</th><th>{h(_("Description"))}</th><th>{h(_("HSN / SAC"))}</th><th>{h(_("Qty."))}</th><th>{h(_("Rate"))}</th><th>{h(_("Amount (INR)"))}</th></tr></thead><tbody>{items}{summary_rows}<tr class="grand"><td colspan="5" class="summary-label">{h(_("Grand Total"))}</td><td class="num">{_money_text(data["grand_total"])}</td></tr></tbody></table>
<div class="declaration"><b>{h(_("Amount in words"))}:</b> {h(data["amount_in_words"])}</div>
{footer_html}
</body></html>"""


def _build_pdf(data):
	"""Use Frappe's PDF engine, with a local fallback when wkhtmltopdf is absent."""
	from frappe.utils.pdf import get_pdf

	html = _render_pdf_html(data)
	try:
		return get_pdf(
			html,
			options={
				"page-size": "A4",
				"margin-top": "8mm",
				"margin-right": "8mm",
				"margin-bottom": "10mm",
				"margin-left": "8mm",
				"encoding": "UTF-8",
			},
		)
	except OSError as error:
		if "wkhtmltopdf" not in str(error).lower():
			raise
		try:
			from weasyprint import HTML
		except ImportError:
			frappe.throw(_("PDF generation is unavailable because no supported PDF renderer is installed."))
		return HTML(string=html).write_pdf()


def _build_docx(data):
	from docx import Document
	from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
	from docx.enum.text import WD_ALIGN_PARAGRAPH
	from docx.oxml import OxmlElement
	from docx.shared import Mm, Pt, RGBColor

	theme = _invoice_style(data)
	doc = Document()
	section = doc.sections[0]
	section.page_width = Mm(210)
	section.page_height = Mm(297)
	section.top_margin = Mm(10)
	section.bottom_margin = Mm(12)
	section.left_margin = Mm(10)
	section.right_margin = Mm(10)
	styles = doc.styles
	styles["Normal"].font.name = theme["font"]
	styles["Normal"].font.size = Pt(9)

	title = doc.add_paragraph()
	title.alignment = {
		"left": WD_ALIGN_PARAGRAPH.LEFT,
		"center": WD_ALIGN_PARAGRAPH.CENTER,
		"right": WD_ALIGN_PARAGRAPH.RIGHT,
	}[theme["align"]]
	run = title.add_run(data["title"])
	run.bold = True
	run.font.size = Pt(16)
	run.font.color.rgb = RGBColor.from_string(theme["accent"])
	copy = doc.add_paragraph(data["copy_label"])
	copy.alignment = WD_ALIGN_PARAGRAPH.RIGHT
	copy.paragraph_format.space_after = Pt(3)

	header = doc.add_table(rows=2, cols=2)
	header.alignment = WD_TABLE_ALIGNMENT.CENTER
	header.style = "Table Grid"
	_set_table_widths(header, (95, 95))
	_add_party_cell(header.cell(0, 0), _("Supplier Details"), data["supplier"], data["invoice_type"])
	header_lines = [
		(_("Invoice No."), data["invoice_number"]),
		(_("Invoice Date"), data["invoice_date"]),
		(_("Buyer's Order No."), data["buyer_order_number"]),
		(_("Date"), data["buyer_order_date"]),
		(_("Supplier's Reference"), data["supplier_reference"]),
		(_("Despatch Document No."), data["dispatch_document_number"]),
		(_("Delivery Note Date"), data["delivery_note_date"]),
	]
	_add_lines(header.cell(0, 1), header_lines, include_empty=True)
	_add_party_cell(header.cell(1, 0), _("Consignee's Details"), data["consignee"], "GST")
	_add_party_cell(
		header.cell(1, 1),
		_("Buyer's Details (if other than Consignee)"),
		data["buyer"],
		"GST",
	)
	for row in header.rows:
		for cell in row.cells:
			cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP

	doc.add_paragraph()
	table = doc.add_table(rows=1, cols=6)
	table.style = "Table Grid"
	table.alignment = WD_TABLE_ALIGNMENT.CENTER
	_set_table_widths(table, (11, 80, 23, 15, 27, 34))
	repeat_header = OxmlElement("w:tblHeader")
	table.rows[0]._tr.get_or_add_trPr().append(repeat_header)
	for cell, text in zip(
		table.rows[0].cells,
		[
			_("Sr. No."),
			_("Description"),
			_("HSN / SAC"),
			_("Qty."),
			_("Rate"),
			_("Amount (INR)"),
		],
		strict=True,
	):
		cell.text = str(text)
		_set_cell_shading(cell, theme["accent"] if theme["frame"] == "band" else theme["soft"])
		for run in cell.paragraphs[0].runs:
			run.bold = True
			if theme["frame"] == "band":
				run.font.color.rgb = RGBColor(255, 255, 255)
	for item in data["items"]:
		cells = table.add_row().cells
		values = [
			item["number"],
			item["description"],
			item["hsn_sac"],
			_quantity_text(item["quantity"]),
			_money_text(item["rate"]),
			_money_text(item["amount"]),
		]
		for index, (cell, value) in enumerate(zip(cells, values, strict=True)):
			cell.text = str(value)
			if index >= 3:
				cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT

	summary = [
		(_("Total"), data["items_total"]),
	]
	if data["invoice_type"] == "GST":
		summary.append((_("GST"), data["gst_amount"]))
	summary.extend(
		[
			(_("Transportation"), data["transportation_charges"]),
			(_("Others"), data["other_charges"]),
		]
	)
	for label, amount in summary:
		_add_summary_row(table, label, amount, bold=False)
	_add_summary_row(table, _("Grand Total"), data["grand_total"], bold=True)

	p = doc.add_paragraph()
	p.add_run(f"{_('Amount in words')}: ").bold = True
	p.add_run(data["amount_in_words"])

	footer_tables = []
	if _has_bank(data["bank"]):
		bank_table = doc.add_table(rows=1, cols=1)
		bank_table.style = "Table Grid"
		bank_table.alignment = WD_TABLE_ALIGNMENT.CENTER
		_set_table_widths(bank_table, (190,))
		bank_cell = bank_table.cell(0, 0)
		bank_cell.text = str(_("Supplier Remittance Details"))
		bank_cell.paragraphs[0].runs[0].bold = True
		_add_lines(bank_cell, _bank_lines(data["bank"]), include_empty=False)
		footer_tables.append(bank_table)

	def add_signature_section(title, statement, signature_png, name, secondary, label):
		signature_table = doc.add_table(rows=1, cols=1)
		signature_table.style = "Table Grid"
		signature_table.alignment = WD_TABLE_ALIGNMENT.CENTER
		_set_table_widths(signature_table, (190,))
		cell = signature_table.cell(0, 0)
		cell.paragraphs[0].add_run(title).bold = True
		if statement:
			cell.add_paragraph(statement)
		p = cell.add_paragraph()
		p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
		p.add_run().add_picture(BytesIO(signature_png), width=Mm(55))
		for value in (name, secondary, label):
			if value:
				p = cell.add_paragraph(value)
				p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
		footer_tables.append(signature_table)

	add_signature_section(
		_("Volunteer reimbursement declaration"),
		data["volunteer_declaration"],
		data["volunteer_signature_png"],
		data["volunteer_name"],
		data["volunteer_employee"],
		_("Volunteer Signature"),
	)
	if data["vendor_will_sign"]:
		add_signature_section(
			_("Vendor confirmation"),
			data["vendor_declaration"],
			data["vendor_signature_png"],
			data["authorised_signatory"] or data["supplier"]["name"],
			_("For {0}").format(data["supplier"]["name"]),
			_("Authorised Signatory"),
		)
	for current_table in (header, table, *footer_tables):
		for row in current_table.rows:
			no_split = OxmlElement("w:cantSplit")
			row._tr.get_or_add_trPr().append(no_split)

	buffer = BytesIO()
	doc.save(buffer)
	return buffer.getvalue()


def _set_table_widths(table, widths_mm):
	from docx.shared import Mm

	table.autofit = False
	for column, width in zip(table.columns, widths_mm, strict=True):
		column.width = Mm(width)
	for row in table.rows:
		for cell, width in zip(row.cells, widths_mm, strict=True):
			cell.width = Mm(width)


def _set_cell_shading(cell, fill):
	from docx.oxml import OxmlElement

	shading = OxmlElement("w:shd")
	shading.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}fill", fill)
	cell._tc.get_or_add_tcPr().append(shading)


def _add_party_cell(cell, heading, party, invoice_type):
	cell.text = ""
	p = cell.paragraphs[0]
	p.add_run(str(heading)).bold = True
	p = cell.add_paragraph(party["name"])
	p.runs[0].bold = True
	for value in (party["address"], " ".join(filter(None, [party["state"], party["pin_code"]]))):
		if value:
			cell.add_paragraph(value)
	identity_label = _("GSTIN / UID No.") if invoice_type == "GST" else _("PAN No.")
	identity = party["gstin"] if invoice_type == "GST" else party["pan"]
	p = cell.add_paragraph()
	p.add_run(f"{identity_label}: ").bold = True
	p.add_run(identity)


def _add_lines(cell, lines, include_empty=False):
	for label, value in lines:
		if not value and not include_empty:
			continue
		p = cell.add_paragraph()
		p.add_run(f"{label}: ").bold = True
		p.add_run(str(value))


def _add_summary_row(table, label, amount, bold):
	cells = table.add_row().cells
	merged = cells[0].merge(cells[4])
	merged.text = str(label)
	merged.paragraphs[0].alignment = 2
	cells[5].text = _money_text(amount)
	cells[5].paragraphs[0].alignment = 2
	if bold:
		for cell in (merged, cells[5]):
			for run in cell.paragraphs[0].runs:
				run.bold = True
